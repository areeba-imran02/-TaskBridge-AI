"""Real spreadsheet integration: pandas + openpyxl. Executes a validated plan."""
import io

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference

AGGS = {"sum", "mean", "count", "min", "max", "median"}


def load_dataframe(uploaded_file) -> pd.DataFrame:
    name = uploaded_file.name.lower()
    if name.endswith(".csv"):
        return pd.read_csv(uploaded_file)
    return pd.read_excel(uploaded_file)


def describe(df: pd.DataFrame) -> str:
    cols = ", ".join(f"{c} ({t})" for c, t in df.dtypes.astype(str).items())
    return f"Rows: {len(df)}\nColumns: {cols}\nSample rows:\n{df.head(5).to_csv(index=False)}"


def _match_col(df, name):
    if name in df.columns:
        return name
    low = {str(c).strip().lower(): c for c in df.columns}
    return low.get(str(name).strip().lower())


def validate(df: pd.DataFrame, params: dict):
    """Returns (errors, normalized_params)."""
    errors = []
    p = {
        "clean": {"drop_duplicates": False, "drop_empty_rows": False, "strip_text": False},
        "totals": [], "group_summaries": [], "charts": [],
    }
    p["clean"].update({k: bool(v) for k, v in (params.get("clean") or {}).items() if k in p["clean"]})

    for t in params.get("totals") or []:
        col, agg = _match_col(df, t.get("column")), t.get("agg", "sum")
        if col is None:
            errors.append(f"Column '{t.get('column')}' was not found in the file.")
        elif agg not in AGGS:
            errors.append(f"Aggregation '{agg}' is not supported.")
        else:
            p["totals"].append({"column": col, "agg": agg})

    for g in params.get("group_summaries") or []:
        gb, col, agg = _match_col(df, g.get("group_by")), _match_col(df, g.get("column")), g.get("agg", "sum")
        if gb is None or col is None:
            errors.append(f"Group summary columns not found: {g.get('group_by')} / {g.get('column')}")
        elif agg not in AGGS:
            errors.append(f"Aggregation '{agg}' is not supported.")
        else:
            p["group_summaries"].append({"group_by": gb, "column": col, "agg": agg})

    for c in params.get("charts") or []:
        idx, typ = c.get("summary_index"), c.get("type", "bar")
        if not isinstance(idx, int) or not (0 <= idx < len(p["group_summaries"])):
            errors.append("Each chart needs a valid group summary.")
        elif typ not in {"bar", "line", "pie"}:
            errors.append(f"Chart type '{typ}' is not supported.")
        else:
            p["charts"].append({"summary_index": idx, "type": typ, "title": c.get("title") or "Chart"})

    if not (p["totals"] or p["group_summaries"] or any(p["clean"].values())):
        errors.append("The plan has no executable operations.")
    return errors, p


def _agg(df, col, agg):
    if agg == "count":
        return df[col].count()
    series = pd.to_numeric(df[col], errors="coerce")
    return getattr(series, agg)()


def execute(df: pd.DataFrame, p: dict):
    """Returns (xlsx_bytes, log_lines, preview_tables)."""
    log, previews = [], {}
    work = df.copy()
    n0 = len(work)

    if p["clean"]["strip_text"]:
        for c in work.select_dtypes(include="object").columns:
            work[c] = work[c].astype(str).str.strip().where(work[c].notna(), None)
        log.append("OK: Trimmed whitespace in text columns.")
    if p["clean"]["drop_empty_rows"]:
        work = work.dropna(how="all")
        log.append(f"OK: Removed {n0 - len(work)} empty rows.")
    if p["clean"]["drop_duplicates"]:
        before = len(work)
        work = work.drop_duplicates()
        log.append(f"OK: Removed {before - len(work)} duplicate rows.")

    totals_rows = []
    for t in p["totals"]:
        val = _agg(work, t["column"], t["agg"])
        totals_rows.append({"Column": t["column"], "Operation": t["agg"], "Result": val})
        log.append(f"OK: {t['agg']}({t['column']}) = {val}")
    totals_df = pd.DataFrame(totals_rows)

    summaries = []
    for i, g in enumerate(p["group_summaries"]):
        tmp = work[[g["group_by"], g["column"]]].copy()
        if g["agg"] != "count":
            tmp[g["column"]] = pd.to_numeric(tmp[g["column"]], errors="coerce")
        s = tmp.groupby(g["group_by"], dropna=False)[g["column"]].agg(g["agg"]).reset_index()
        s.columns = [g["group_by"], f"{g['agg']}_{g['column']}"]
        summaries.append(s)
        previews[f"Summary {i + 1}"] = s
        log.append(f"OK: {g['agg']} of {g['column']} by {g['group_by']} ({len(s)} groups).")

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        work.to_excel(xw, sheet_name="Cleaned Data", index=False)
        if len(totals_df):
            totals_df.to_excel(xw, sheet_name="Totals", index=False)
        for i, s in enumerate(summaries):
            s.to_excel(xw, sheet_name=f"Summary {i + 1}", index=False)
        for c in p["charts"]:
            ws = xw.book[f"Summary {c['summary_index'] + 1}"]
            n = len(summaries[c["summary_index"]])
            chart = {"bar": BarChart, "line": LineChart, "pie": PieChart}[c["type"]]()
            chart.title = c["title"]
            data = Reference(ws, min_col=2, min_row=1, max_row=n + 1)
            cats = Reference(ws, min_col=1, min_row=2, max_row=n + 1)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            ws.add_chart(chart, "E2")
            log.append(f"OK: {c['type']} chart '{c['title']}' created.")

    data_bytes = buf.getvalue()

    # Evidence check: re-open the generated file and verify it
    wb = load_workbook(io.BytesIO(data_bytes))
    expected = 1 + (1 if len(totals_df) else 0) + len(summaries)
    charts_found = sum(len(ws._charts) for ws in wb.worksheets)
    if len(wb.sheetnames) != expected or charts_found != len(p["charts"]):
        raise RuntimeError("Verification failed: the output workbook does not match the expected structure.")
    log.append(f"VERIFIED: Re-opened the workbook and found {len(wb.sheetnames)} sheets and {charts_found} charts.")

    previews = {"Cleaned Data": work.head(50), **({"Totals": totals_df} if len(totals_df) else {}), **previews}
    return data_bytes, log, previews
